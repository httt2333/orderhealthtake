"""Build a submission document from supplied images and explicit UI mockups.
Run with the bundled workspace Python runtime. No model inference is performed.
"""
from pathlib import Path
import json
from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'artifacts' / 'recognition-report'
OUT.mkdir(parents=True, exist_ok=True)
CASES = json.loads((ROOT/'assets/demo-packages/catalog.json').read_text(encoding='utf-8'))
FONT = 'C:/Windows/Fonts/msyh.ttc'
BOLD = 'C:/Windows/Fonts/msyhbd.ttc'
GREEN, INK, BG = '#236d53', '#183c36', '#faf9f6'
S = 2

def f(size, bold=False):
    return ImageFont.truetype(BOLD if bold else FONT, round(size*S))

def txt(draw, text, xy, width=286, size=21, bold=False, fill=INK, gap=5):
    x,y=xy;line='';font=f(size,bold)
    for ch in text:
        if ch=='\n' or (line and draw.textlength(line+ch,font=font)>width*S):
            draw.text((x*S,y*S),line,font=font,fill=fill);y+=size+gap;line='' if ch=='\n' else ch
        else: line+=ch
    if line: draw.text((x*S,y*S),line,font=font,fill=fill);y+=size+gap
    return y

def rect(draw, box, fill, radius=14, outline=None, width=1):
    draw.rounded_rectangle(tuple(round(v*S) for v in box),radius=radius*S,fill=fill,outline=outline,width=width*S)

def picture(canvas, case, box):
    x,y,w,h=box
    photo=Image.open(ROOT/CASES[case]['imagePath']).convert('RGB')
    photo=ImageOps.contain(photo,(int(w*S),int(h*S)))
    canvas.paste(photo,(int((x+w/2)*S-photo.width/2),int((y+h/2)*S-photo.height/2)))

def button(draw,text,y,primary=True):
    rect(draw,(20,y,320,y+52),GREEN if primary else '#edf4ef')
    font=f(21,True);w=draw.textlength(text,font=font)
    draw.text(((340*S-w)/2,(y+12)*S),text,font=font,fill='white' if primary else INK)

def field(draw,label,value,y):
    txt(draw,label,(24,y),size=19,bold=True)
    rect(draw,(22,y+29,318,y+75),'white',outline='#c5d9cc')
    txt(draw,value,(34,y+39),width=268,size=20,fill='#456552')
    return y+91

def phone(key,title,mode,case=0):
    im=Image.new('RGB',(340*S,710*S),'white');d=ImageDraw.Draw(im)
    rect(d,(1,1,339,709),BG,22,'#b6cbbd',2)
    txt(d,'安心药盒',(20,20),size=17,bold=True)
    txt(d,'页面示意',(247,23),width=78,size=12,fill='#63766a')
    txt(d,title,(20,67),size=26,bold=True)
    if mode=='start':
        rect(d,(52,151,288,368),'#e5efe7',outline='#b4cdbd')
        d.rectangle((112*S,202*S,228*S,297*S),outline=GREEN,width=5*S)
        d.ellipse((148*S,219*S,193*S,266*S),outline=GREEN,width=5*S)
        txt(d,'对准包装正面',(79,324),size=23,bold=True)
        button(d,'拍药盒',408);button(d,'看识别示例',480,False)
    elif mode=='choose':
        for i in range(3):
            y=143+i*141;rect(d,(20,y,320,y+127),'white',outline='#d6e3da')
            picture(im,i,(28,y+8,99,108))
            txt(d,CASES[i]['label'],(143,y+26),width=155,size=21,bold=True)
        txt(d,'其余两个包装继续向下看',(30,580),size=16)
    elif mode=='recognize':
        txt(d,'示例识别结果',(20,111),size=16,fill='#63766a')
        picture(im,case,(24,147,292,227))
        y=txt(d,CASES[case]['name'],(24,391),size=24,bold=True)
        y=txt(d,CASES[case]['spec'],(24,y+5),size=17,width=290)
        txt(d,CASES[case]['note'],(24,min(y+9,495)),size=16,width=290,fill='#7b5423')
        button(d,'是这个包装，继续',555)
    elif mode=='form':
        picture(im,case,(24,133,292,127))
        y=field(d,'药品名称',CASES[case]['name'],276)
        y=field(d,'每次数量','按药单填写',y)
        field(d,'服用时间','-- : --',y)
        button(d,'下一步核对',555)
    elif mode in ['confirm','today','read','record','large']:
        ph=306 if mode=='large' else 247
        picture(im,case,(22,132,296,ph))
        y=132+ph+10
        y=txt(d,CASES[case]['name'],(24,y),size=22,bold=True)
        y=txt(d,'每次  已核对数量',(24,y+6),size=21,bold=True)
        txt(d,'时间  已核对时间',(24,y+5),size=19)
        if mode=='confirm':button(d,'听一听',491,False);button(d,'确认并保存',555)
        elif mode=='read':button(d,'正在朗读  点此重听',491);button(d,'停止',555,False)
        elif mode=='record':button(d,'已记录服用',491,False);button(d,'撤销记录',555,False)
        elif mode=='large':button(d,'返回今天用药',555,False)
        else:button(d,'听这次怎么吃',491);button(d,'我已服用     记不清',555,False)
    # The fixed voice entry is the same across the designed pages.
    d.rectangle((2*S,622*S,338*S,708*S),fill=BG)
    txt(d,'点我，说说您想做什么',(76,631),size=16,bold=True)
    d.ellipse((145*S,659*S,195*S,709*S),fill=GREEN)
    d.rounded_rectangle((164*S,669*S,176*S,690*S),radius=6*S,outline='white',width=2*S)
    d.arc((159*S,674*S,181*S,695*S),0,180,fill='white',width=2*S)
    d.line((170*S,695*S,170*S,701*S),fill='white',width=2*S)
    target=OUT/(key+'.png');im.save(target);return target

MOCKS={
 'start':phone('01-start','拍下这个药盒','start'),
 'choose':phone('02-choose','选择药盒示例','choose'),
 'recognize':phone('03-recognize','是这个包装吗','recognize'),
 'form':phone('04-form','填写用药安排','form'),
 'confirm':phone('05-confirm','照片和安排对吗','confirm'),
 'today':phone('06-today','看药盒  听安排','today'),
 'large':phone('07-large','看清这个药盒','large'),
 'read':phone('08-read','听这次怎么吃','read'),
 'record':phone('09-record','今天用药','record'),
 'red':phone('10-red','是这个包装吗','recognize',2),
 'orange':phone('11-orange','是这个包装吗','recognize',3),
 'child':phone('12-child','先核对使用人','recognize',4),
}

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.6);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)
for style_name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[style_name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(20)
doc.styles['Heading 2'].font.size=Pt(13)
doc.styles['Caption'].font.size=Pt(9)
for style_name in ['Title','Heading 1','Heading 2','Normal']:
    doc.styles[style_name].paragraph_format.page_break_before=False
    for border in doc.styles[style_name].element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
header=sec.header.paragraphs[0];header.text='安心药盒  场景 A  产品演示材料';header.style='Caption'
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('包装图片由提交者提供  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def p(text,bold=False):
    para=doc.add_paragraph();r=para.add_run(text);r.bold=bold;return para
def h(text):doc.add_heading(text,1)
def newpage(title):doc.add_page_break();h(title)
def phones(keys,labels):
    para=doc.add_paragraph();para.paragraph_format.space_after=Pt(2)
    for i,key in enumerate(keys):
        para.add_run().add_picture(str(MOCKS[key]),width=Inches(2.12))
        if i<len(keys)-1:para.add_run(' ')
    cap=doc.add_paragraph('    '.join(labels),style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
def table(headers,rows,widths):
    tb=doc.add_table(rows=1,cols=len(headers));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    for i,width in enumerate(widths):tb.columns[i].width=Inches(width)
    for i,title in enumerate(headers):tb.rows[0].cells[i].text=title
    for values in rows:
        cells=tb.add_row().cells
        for i,value in enumerate(values):cells[i].text=str(value)
    for ri,row in enumerate(tb.rows):
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'5');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
            tcPr.append(borders);margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'100');e.set(qn('w:type'),'dxa');margins.append(e)
            tcPr.append(margins);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'294D61' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF'));tcPr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.15
                for run in para.runs:
                    run.font.size=Pt(10);run.bold=ri==0
                    run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
        if ri==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return tb

doc.add_paragraph('安心药盒 包装识别演示与页面方案',style='Title')
p('以五张实际包装图片说明拍照后如何核对信息、填写既定安排，并通过药盒大图和语音完成每日用药任务。提交场景为能独立生活、长期使用多种药物的老人。')
p('当前版本实现了图片示例、预置结果核对、手动安排录入、照片绑定和逐项朗读。五例是独立演示素材，不构成一份共同服用方案。识别结果为依据图片整理的预置数据，不是接入真实 AI 后获得的实测结果。',True)
doc.add_heading('从拍照到核对',2)
phones(['start','choose','recognize'],['页面一  拍摄入口','页面二  五个独立示例','页面三  包装核对'])
p('应用中只呈现照片、名称、包装关键字和一个下一步动作。详细来源、识别限制和验收条件放在本材料中，不堆到老人操作页面。')
p('以下手机图均为页面设计示意，不是实机截图；其中带“已核对数量／时间”的位置是设计占位字段，实际空值不允许保存。')

newpage('五张输入图片与预期读取字段')
strip=Image.new('RGB',(1500,370),'white')
for i,c in enumerate(CASES):
    im=Image.open(ROOT/c['imagePath']).convert('RGB');im=ImageOps.contain(im,(280,290));strip.paste(im,(i*300+(300-im.width)//2,(300-im.height)//2))
    ImageDraw.Draw(strip).text((i*300+112,318),f'案例 {i+1}',font=ImageFont.truetype(FONT,24),fill=INK)
strip_path=OUT/'evidence-strip.png';strip.save(strip_path)
doc.add_picture(str(strip_path),width=Inches(7.05))
table(['案例','可读名称','包装可见信息','需要核对'],[
 ['1','阿莫西林胶囊','规格 0.25g\n每小盒24粒','规格与包装粒数均不等于每次用量'],
 ['2','头孢地尼胶囊','0.1g × 6粒','拍摄小字偏模糊，最终以实物核对'],
 ['3','维生素C含片','红白瓶\n营养素补充剂\n净含量60g','品牌、完整小字和单片规格不自动补全'],
 ['4','维生素C含片','品健橙色包装\n65g（100片）\n保健食品不是药物','与红白瓶分别建立身份，不按同名合并'],
 ['5','维生素C含片','包装正面含“儿童”标记','名称后缀不完整，核对完整名称和使用人'],
],[.55,1.4,2.2,2.85])
p('这里记录的是图片上可见的包装文字。没有核实真伪、适用人群、适应症、相互作用或处方正确性；也没有从包装上的用法小字生成个人服用安排。')
p('两种胶囊、两款一般维 C 包装及儿童标记包装用于覆盖不同识别情况。它们只是独立测试输入，并不代表适合老人共同使用。',True)

newpage('核对后如何进入每日任务')
p('确认包装后，仅把对应照片与名称带入录入页。每次数量与时间保持空白，用户依据既有药单填写。保存前再次核对照片、药名、数量和时间。')
phones(['form','confirm','today'],['页面四  数量时间不预填','页面五  对照实物和药单','页面六  看照片并听安排'])
doc.add_heading('页面状态与保存条件',2)
table(['页面','主要内容','继续条件'],[
 ['填写安排','原照片与名称可编辑；数量、时间为空','照片可用且三项内容完整'],
 ['保存核对','同一照片、名称、已填写数量与时间；可试听','用户确认后才写入今日安排'],
 ['今天用药','大照片与逐项朗读在前，记录服药在后','示例来源持续标为演示安排'],
],[1.0,3.6,2.4])
p('“已核对数量”和“已核对时间”是本材料的版式占位，不能作为真实安排写入产品。图片上的0.25g、0.1g或65g也不能用来替代这些字段。')

newpage('看清照片与听清安排')
p('药名作为辅助信息。老人首先对照自家药盒外观，照片不清时点开放大；听不清时重听或停止。每段朗读只读取当前这一项已经核对的名称、数量和时间。')
phones(['large','read','record'],['页面七  放大后仍是同一药盒','页面八  朗读反馈与停止','页面九  服药后的记录'])
doc.add_heading('常驻语音入口',2)
p('每页下方保留圆形麦克风。默认提示“点我，说说您想做什么”；点击才开始收音，正在听时保留麦克风并出现停止标记。任务识别成功先回应要打开的页面，再进入相应任务。')
doc.add_heading('避免把未知情况写成完成',2)
p('“我已服用”是本人记录，不是照片或传感器证明。记不清是否服过时单独保留不确定状态，提示查看药单或咨询医生、药师；产品不生成补服方案。')

newpage('同名包装与儿童标记的核对')
p('识别不能只返回一个通用名称。维生素C含片可能有不同品牌、外观和包装文字，示例必须保持照片与条目一一对应；包装带儿童标记时，应把核对使用人放在继续录入之前。')
phones(['red','orange','child'],['红白瓶  小字不强行补全','品健包装  保留包装声明','儿童标记  先核对使用人'])
table(['异常或不确定','简短的页面处理','后续动作'],[
 ['同名不同包装','仍显示各自照片与包装信息','用户分别核对，不能静默合并'],
 ['儿童标记','请核对完整名称和使用人','核对包装后再手动填写既定安排'],
 ['图片模糊或字段缺失','重新拍摄／自己填写','实际AI的模糊拒识与缺字段能力仍待接入测试'],
],[1.5,2.7,2.8])

newpage('验证记录与提交说明')
table(['验证层次','本轮得到什么','尚不能声称什么'],[
 ['图片整理','依据五张用户图片人工转录可读名称与包装字样','不是药品真伪或医疗适用性核验'],
 ['识别演示','五个预置结果与原图绑定，完整流程可运行','不是模型识别准确率或置信度实测'],
 ['代码回归','检查独立身份、空数量时间、同名不合并、儿童提示和未知例拒绝','不证明真实相机、中文音色或老人独立可用性'],
 ['真实AI评测','预留同一组输入与预期字段供接入后比较','当前没有模型输出，不能报告命中率'],
],[1.1,3.0,2.9])
doc.add_heading('真实模型接入后的评测方法',2)
p('用原图、模糊图、局部图与未收录包装进行测试。逐项核对名称、规格、包装数量及儿童标记；记录模型是否遗漏、读错或拒绝。特别检查它是否把规格误填为每次数量。建立人工核对后的字段基准，再报告结果，不只统计“认出了药名”。')
doc.add_heading('运行与能力说明',2)
p('Node.js 18及以上，在项目目录运行 npm run dev，打开 http://localhost:4173/#/elder/setup，选择“试用示例”。五张图已在项目内，无需评审自行上传。真实拍摄入口当前只保存照片并允许手动填写，不执行实时AI识别。')
p('逐项朗读、照片绑定与本机保存已经实现。图片识别示例为预置结果；自动OCR、生成式AI、后台推送和跨设备共享未实现。关闭或隐藏网页后不保证提醒。')
doc.add_heading('素材与工具声明',2)
p('图片来源为提交者在本轮提供的五张包装图片，顺序对应案例1至5。来源名称保存在 assets/demo-packages/catalog.json。本材料未引用外部医学或市场研究结论。')
p('使用工具为 Codex 辅助产品方案、编码、测试与文档制作；浏览器提供语音识别和语音合成。实际投入时间由提交者按真实记录补充，不能用自动生成内容代填。')
doc.core_properties.title='安心药盒 包装识别演示与页面方案'
doc.core_properties.subject='场景A 五张包装图片的预置识别演示与页面状态说明'
doc.core_properties.author='安心药盒项目'
target=ROOT/'挑战交付/安心药盒_包装识别演示与页面方案.docx'
doc.save(target)
print(target)
